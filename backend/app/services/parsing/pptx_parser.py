# app/services/parsing/pptx_parser.py
# PowerPoint parser — Section 5.1
from __future__ import annotations
import io
import logging
from typing import List

from app.services.parsing.pdf import ParsedItem

logger = logging.getLogger(__name__)


class PPTXParser:
    """Extract text, titles, notes from PowerPoint slides. One ParsedItem per slide."""

    def parse(self, file_bytes: bytes) -> List[ParsedItem]:
        if not file_bytes:
            raise ValueError("Empty file bytes provided to PPTXParser")

        try:
            from pptx import Presentation
        except ImportError as e:
            raise RuntimeError(
                "python-pptx is not installed. Run: pip install python-pptx"
            ) from e

        try:
            prs = Presentation(io.BytesIO(file_bytes))
        except Exception as e:
            raise RuntimeError(
                f"Failed to open PowerPoint file — it may be corrupted or in an old .ppt format "
                f"(only .pptx is supported). Error: {e}"
            ) from e

        items: List[ParsedItem] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            title_text = ""
            body_parts: list[str] = []

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                try:
                    text = shape.text_frame.text.strip()
                except Exception:
                    continue

                if not text:
                    continue

                # Detect title placeholder (idx 0 = title, idx 1 = body)
                is_title = False
                try:
                    if (
                        shape.name.lower().startswith("title")
                        or (
                            hasattr(shape, "placeholder_format")
                            and shape.placeholder_format
                            and shape.placeholder_format.idx == 0
                        )
                    ):
                        is_title = True
                except Exception:
                    pass

                if is_title:
                    title_text = text
                else:
                    body_parts.append(text)

            # Speaker notes
            notes_text = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                pass

            body = "\n".join(body_parts)
            if notes_text:
                body += f"\n\nNotes: {notes_text}"

            # Skip completely blank slides
            full_text = f"Slide {slide_num}: {title_text}\n{body}".strip()
            if not full_text or full_text == f"Slide {slide_num}:":
                logger.debug(f"[PPTXParser] Slide {slide_num} is blank — skipping")
                continue

            items.append(
                ParsedItem(
                    page_num=slide_num,
                    slide_num=slide_num,
                    content_type="text",
                    text=full_text,
                    title=title_text,
                    body=body,
                )
            )

        logger.info(f"[PPTXParser] Extracted {len(items)} slides from {len(prs.slides)}-slide presentation")

        if not items:
            raise RuntimeError(
                "No text content found in the PowerPoint file. "
                "The slides may contain only images or the file may be empty."
            )

        return items
