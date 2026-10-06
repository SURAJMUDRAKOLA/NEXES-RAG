# app/services/parsing/docx_parser.py
# Word document parser — plan §4.1
from __future__ import annotations
import io
import logging
from typing import List

from app.services.parsing.pdf import ParsedItem

logger = logging.getLogger(__name__)


class DOCXParser:
    """Extract text and tables from Word documents (.docx)."""

    def parse(self, file_bytes: bytes) -> List[ParsedItem]:
        if not file_bytes:
            raise ValueError("Empty file bytes")
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            items: List[ParsedItem] = []
            page_num = 1
            para_buffer: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    para_buffer.append(text)
                    if len(para_buffer) >= 50:
                        items.append(ParsedItem(
                            page_num=page_num,
                            content_type="text",
                            text="\n".join(para_buffer),
                        ))
                        page_num += 1
                        para_buffer = []

            if para_buffer:
                items.append(ParsedItem(
                    page_num=page_num,
                    content_type="text",
                    text="\n".join(para_buffer),
                ))

            for table in doc.tables:
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                md = "\n".join(rows)
                if md.strip():
                    items.append(ParsedItem(
                        page_num=page_num,
                        content_type="table",
                        text=md,
                        markdown_text=md,
                    ))

            if not items:
                raise RuntimeError("No content found in Word document")
            logger.info(f"[DOCXParser] Extracted {len(items)} items")
            return items
        except Exception as e:
            raise RuntimeError(f"DOCX parse failed: {e}") from e
